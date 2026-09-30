"""
Motor unificado de importación masiva de catálogo Keyestudio (Ajustes #10, #11, #12, #20 y #21 de Humm).
Servicio reutilizado idénticamente por la CLI (importar_catalogo_keyestudio) y por la interfaz web (/gestion/importaciones/).
"""

import os
import re
import shutil
import zipfile
from collections import defaultdict
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import transaction
from PIL import Image
import openpyxl

from apps.catalogo.models import Proveedor, Categoria, Producto, ProductoImagen
from apps.core.models import ConfiguracionPricing
from apps.gestion.services_auditoria import registrar_actividad


def parse_precio_usd(val):
    """Normaliza valores de texto o numéricos a Decimal de dos dígitos."""
    if val is None:
        return None
    if isinstance(val, (int, float, Decimal)):
        try:
            return Decimal(str(val)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        except (InvalidOperation, ValueError):
            return None

    val_str = str(val).strip()
    if not val_str:
        return None

    val_clean = re.sub(r"[^\d,\.]", "", val_str)
    if not val_clean:
        return None

    if "," in val_clean and "." in val_clean:
        if val_clean.rfind(".") > val_clean.rfind(","):
            val_clean = val_clean.replace(",", "")
        else:
            val_clean = val_clean.replace(".", "").replace(",", ".")
    elif "," in val_clean:
        val_clean = val_clean.replace(",", ".")

    try:
        return Decimal(val_clean).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    except (InvalidOperation, ValueError):
        return None


def optimizar_imagen(source_path, dest_path, max_width=1200, quality=85):
    """Copia y optimiza una imagen fuente con Pillow sin destruir el archivo original."""
    dest_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        with Image.open(source_path) as img:
            if img.mode in ("RGBA", "P") and dest_path.suffix.lower() in (".jpg", ".jpeg"):
                img = img.convert("RGB")
            
            if img.width > max_width:
                height = int(img.height * (max_width / img.width))
                img = img.resize((max_width, height), Image.Resampling.LANCZOS)

            save_kwargs = {}
            if dest_path.suffix.lower() in (".jpg", ".jpeg"):
                save_kwargs = {"quality": quality, "optimize": True}
            elif dest_path.suffix.lower() == ".webp":
                save_kwargs = {"quality": quality}
            
            img.save(dest_path, **save_kwargs)
            return True
    except Exception:
        shutil.copy2(source_path, dest_path)
        return True


class ImportacionCatalogoService:
    # Límites de seguridad para archivos subidos (Ajuste Obligatorio #11)
    MAX_EXCEL_BYTES = 15 * 1024 * 1024       # 15 MB
    MAX_ZIP_BYTES = 50 * 1024 * 1024         # 50 MB
    MAX_ZIP_UNCOMPRESSED = 150 * 1024 * 1024 # 150 MB
    MAX_ZIP_FILES = 2000
    EXTENSIONES_IMAGEN_VALIDAS = {".jpg", ".jpeg", ".png", ".webp"}

    @classmethod
    def validar_archivo_excel(cls, archivo_path):
        """Valida tamaño y formato del archivo Excel."""
        p = Path(archivo_path)
        if not p.exists():
            raise ValidationError(f"El archivo Excel no existe: {p}")
        if p.suffix.lower() not in (".xlsx", ".xls"):
            raise ValidationError("El archivo debe tener extensión .xlsx o .xls.")
        if p.stat().st_size > cls.MAX_EXCEL_BYTES:
            raise ValidationError(f"El archivo Excel excede el tamaño máximo permitido de {cls.MAX_EXCEL_BYTES // (1024*1024)} MB.")
        return True

    @classmethod
    def validar_y_descomprimir_zip(cls, zip_path, destino_dir):
        """
        Valida exhaustivamente un archivo ZIP para prevenir ZIP bombs, path traversal
        y archivos maliciosos antes de extraer (Ajuste Obligatorio #11).
        """
        p = Path(zip_path)
        if not p.exists():
            raise ValidationError("El archivo ZIP no existe.")
        if p.stat().st_size > cls.MAX_ZIP_BYTES:
            raise ValidationError(f"El archivo ZIP excede el tamaño máximo de {cls.MAX_ZIP_BYTES // (1024*1024)} MB.")

        destino = Path(destino_dir)
        destino.mkdir(parents=True, exist_ok=True)

        total_size = 0
        total_files = 0

        with zipfile.ZipFile(p, "r") as zf:
            for member in zf.infolist():
                # Prevenir path traversal (../ o rutas absolutas)
                norm_name = os.path.normpath(member.filename)
                if norm_name.startswith("..") or os.path.isabs(norm_name):
                    raise ValidationError(f"Entrada ZIP sospechosa (path traversal): {member.filename}")

                if member.is_dir():
                    continue

                total_files += 1
                if total_files > cls.MAX_ZIP_FILES:
                    raise ValidationError(f"El archivo ZIP contiene más de {cls.MAX_ZIP_FILES} archivos (límite de seguridad).")

                total_size += member.file_size
                if total_size > cls.MAX_ZIP_UNCOMPRESSED:
                    raise ValidationError(f"El contenido descomprimido excede el límite de {cls.MAX_ZIP_UNCOMPRESSED // (1024*1024)} MB (posible ZIP bomb).")

                # Validar extensión interna
                ext = Path(member.filename).suffix.lower()
                if ext not in cls.EXTENSIONES_IMAGEN_VALIDAS:
                    continue  # Omitir archivos no imagen sin fallar

                # Extracción segura
                target_path = destino / Path(member.filename).name
                with zf.open(member) as source, open(target_path, "wb") as target:
                    shutil.copyfileobj(source, target)

        return destino

    @classmethod
    def procesar_catalogo(cls, excel_path, imagenes_dir=None, is_dry_run=True,
                          actualizar_costos=True, limite=0, usuario=None, lote=None, request=None):
        """
        Ejecuta el procesamiento integral del catálogo:
        - Soporta --dry-run
        - Detecta SKUs únicos y conflictivos
        - Preserva estrictamente la curaduría existente (Ajuste #18)
        - Aplica nuevos productos con publicado=False y SIN_REVISAR (Ajuste #17)
        - Retorna reporte consolidado
        """
        excel_p = Path(excel_path)
        cls.validar_archivo_excel(excel_p)

        imagenes_p = Path(imagenes_dir) if imagenes_dir else None
        if imagenes_p and not imagenes_p.exists():
            imagenes_p = None

        # 1. Parámetros de Pricing
        pricing_config = ConfiguracionPricing.get_solo()
        tc = pricing_config.tipo_cambio_usd_clp
        factor_int = pricing_config.factor_internacion_flete_porcentaje
        recargo_comercial = pricing_config.recargo_general_porcentaje
        iva = pricing_config.iva_porcentaje

        # 2. Indizar banco de imágenes
        banco_imagenes = defaultdict(list)
        imagenes_totales_banco = 0
        if imagenes_p and imagenes_p.is_dir():
            for img_file in imagenes_p.iterdir():
                if img_file.is_file() and img_file.suffix.lower() in cls.EXTENSIONES_IMAGEN_VALIDAS:
                    imagenes_totales_banco += 1
                    stem = img_file.stem.upper().strip()
                    if img_file not in banco_imagenes[stem]:
                        banco_imagenes[stem].append(img_file)
                    tokens = [t.strip() for t in re.split(r"[\s]+|-(?=[A-Za-z0-9]{4,})", stem) if t.strip()]
                    for tok in tokens:
                        if img_file not in banco_imagenes[tok]:
                            banco_imagenes[tok].append(img_file)
                        tok_base = re.split(r"[-_]\d+$", tok)[0]
                        if tok_base and img_file not in banco_imagenes[tok_base]:
                            banco_imagenes[tok_base].append(img_file)

        # 3. Leer Excel
        wb = openpyxl.load_workbook(excel_p, read_only=True, data_only=True)
        sheet = wb.active

        header_row = None
        for row in sheet.iter_rows(max_row=5, values_only=True):
            if any(isinstance(c, str) and ("SKU" in c.upper() or "DESCRIPCI" in c.upper()) for c in row if c):
                header_row = [str(c).strip() if c else "" for c in row]
                break

        if not header_row:
            wb.close()
            raise ValidationError("No se encontró fila de encabezados válida en las primeras 5 filas del Excel.")

        col_map = {}
        for idx, col in enumerate(header_row):
            col_upper = col.upper()
            if "SKU" in col_upper or "ID" in col_upper:
                col_map["sku"] = idx
            elif "DESCRIPCI" in col_upper or "PRODUCTO" in col_upper:
                col_map["descripcion"] = idx
            elif "MINIATURA" in col_upper:
                col_map["miniatura"] = idx
            elif "PRECIO" in col_upper:
                col_map["precio"] = idx
            elif "ALTA RESOLUCI" in col_upper or "IMAGEN ALTA" in col_upper:
                col_map["imagen_hd"] = idx

        if "sku" not in col_map or "descripcion" not in col_map or "precio" not in col_map:
            wb.close()
            raise ValidationError(f"Faltan columnas requeridas (SKU, Descripción, Precio). Encabezados: {header_row}")

        filas_totales_leidas = 0
        filas_sin_sku = 0
        filas_sin_precio = 0
        precios_no_interpretables = 0
        registros_por_sku = defaultdict(list)

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if limite > 0 and filas_totales_leidas >= limite:
                break
            if not any(row):
                continue
            filas_totales_leidas += 1

            raw_sku = row[col_map["sku"]] if col_map["sku"] < len(row) else None
            raw_desc = row[col_map["descripcion"]] if col_map["descripcion"] < len(row) else None
            raw_precio = row[col_map["precio"]] if col_map["precio"] < len(row) else None

            if not raw_sku or not str(raw_sku).strip():
                filas_sin_sku += 1
                continue
            if raw_precio is None or str(raw_precio).strip() == "":
                filas_sin_precio += 1
                continue

            precio_usd = parse_precio_usd(raw_precio)
            if precio_usd is None or precio_usd <= 0:
                precios_no_interpretables += 1
                continue

            sku_norm = str(raw_sku).strip().upper()
            desc_norm = str(raw_desc).strip() if raw_desc else ""

            registros_por_sku[sku_norm].append({
                "sku": sku_norm,
                "descripcion": desc_norm,
                "precio_usd": precio_usd,
            })

        wb.close()

        # 4. Análisis de duplicados y conflictos
        skus_simples = {}
        skus_duplicados_identicos = {}
        skus_conflictivos = {}

        for sku, items in registros_por_sku.items():
            if len(items) == 1:
                skus_simples[sku] = items[0]
            else:
                precios = {it["precio_usd"] for it in items}
                if len(precios) == 1:
                    skus_duplicados_identicos[sku] = items[0]
                else:
                    skus_conflictivos[sku] = items

        candidatos_validos = {**skus_simples, **skus_duplicados_identicos}

        # 5. Comparar contra base de datos existente
        skus_existentes_qs = Producto.objects.filter(sku_proveedor__in=candidatos_validos.keys())
        existentes_map = {p.sku_proveedor: p for p in skus_existentes_qs}

        productos_nuevos_count = 0
        productos_actualizados_costo = 0
        productos_sin_cambios = 0
        conflictos_count = len(skus_conflictivos)
        productos_sin_imagen = 0
        imagenes_asociadas = 0

        precios_usd_validos = [c["precio_usd"] for c in candidatos_validos.values()]
        precio_min = min(precios_usd_validos) if precios_usd_validos else Decimal("0.00")
        precio_max = max(precios_usd_validos) if precios_usd_validos else Decimal("0.00")
        precio_prom = sum(precios_usd_validos) / len(precios_usd_validos) if precios_usd_validos else Decimal("0.00")

        # Proveedor y categoría base
        prov, _ = Proveedor.objects.get_or_create(
            codigo="KEY",
            defaults={"nombre": "Keyestudio", "moneda_origen": "USD"}
        )
        cat_sin_clasificar, _ = Categoria.objects.get_or_create(
            nombre="Sin clasificar",
            defaults={"descripcion": "Componentes pendientes de curaduría pedagógica"}
        )

        # 6. Ejecución definitiva si no es dry-run
        if not is_dry_run:
            with transaction.atomic():
                for sku, data in candidatos_validos.items():
                    prod_existente = existentes_map.get(sku)
                    if prod_existente:
                        # Producto existente: solo actualizar costo USD si cambió (Ajuste #18: Preservación Curaduría)
                        if actualizar_costos and prod_existente.costo_proveedor_usd != data["precio_usd"]:
                            prod_existente.costo_proveedor_usd = data["precio_usd"]
                            prod_existente.save()  # recalcula sugeridos
                            productos_actualizados_costo += 1
                        else:
                            productos_sin_cambios += 1
                        prod_target = prod_existente
                    else:
                        # Producto nuevo: publicado=False, estado_curaduria='SIN_REVISAR' (Ajuste #17)
                        prod_nuevo = Producto(
                            sku_humm=f"HUMM-KEY-{sku}",
                            sku_proveedor=sku,
                            proveedor=prov,
                            categoria=cat_sin_clasificar,
                            marca="Keyestudio",
                            nombre_comercial=data["descripcion"] or f"Componente Keyestudio {sku}",
                            nombre_original_proveedor=data["descripcion"] or "",
                            costo_proveedor_usd=data["precio_usd"],
                            publicado=False,
                            activo=True,
                            estado_curaduria="SIN_REVISAR",
                            estado_especificacion_neutral="NO_REVISADO",
                        )
                        prod_nuevo.save()
                        productos_nuevos_count += 1
                        prod_target = prod_nuevo

                    # Manejo de imágenes
                    archivos_img = banco_imagenes.get(sku, [])
                    if archivos_img:
                        imagenes_asociadas += 1
                        # Si no tiene imágenes, asociar
                        if not prod_target.imagenes.exists():
                            for idx, img_src in enumerate(archivos_img[:5]):
                                dest_media = Path(settings.MEDIA_ROOT) / "catalogo" / "originales" / img_src.name
                                optimizar_imagen(img_src, dest_media)
                                ProductoImagen.objects.create(
                                    producto=prod_target,
                                    archivo=f"catalogo/originales/{img_src.name}",
                                    nombre_archivo_original=img_src.name,
                                    es_principal=(idx == 0),
                                    orden=idx,
                                )
                    else:
                        productos_sin_imagen += 1
        else:
            # En dry-run solo calcular estadísticas
            for sku, data in candidatos_validos.items():
                prod_existente = existentes_map.get(sku)
                if prod_existente:
                    if prod_existente.costo_proveedor_usd != data["precio_usd"]:
                        productos_actualizados_costo += 1
                    else:
                        productos_sin_cambios += 1
                else:
                    productos_nuevos_count += 1

                if banco_imagenes.get(sku):
                    imagenes_asociadas += 1
                else:
                    productos_sin_imagen += 1

        # Resumen estructurado
        resumen = {
            "is_dry_run": is_dry_run,
            "filas_totales_leidas": filas_totales_leidas,
            "filas_sin_sku": filas_sin_sku,
            "filas_sin_precio": filas_sin_precio,
            "precios_no_interpretables": precios_no_interpretables,
            "skus_detectados": len(registros_por_sku),
            "skus_simples": len(skus_simples),
            "skus_duplicados_identicos": len(skus_duplicados_identicos),
            "conflictos_detectados": conflictos_count,
            "productos_candidatos": len(candidatos_validos),
            "productos_nuevos": productos_nuevos_count,
            "productos_actualizados": productos_actualizados_costo,
            "productos_sin_cambios": productos_sin_cambios,
            "precio_min_usd": float(precio_min),
            "precio_max_usd": float(precio_max),
            "precio_prom_usd": float(precio_prom),
            "imagenes_encontradas_banco": imagenes_totales_banco,
            "productos_con_imagen": imagenes_asociadas,
            "productos_sin_imagen": productos_sin_imagen,
            "skus_conflictivos_lista": [
                {"sku": k, "filas": v} for k, v in list(skus_conflictivos.items())[:20]
            ],
        }

        # Actualizar lote en BD si existe
        if lote:
            lote.filas_leidas = filas_totales_leidas
            lote.skus_detectados = len(registros_por_sku)
            lote.productos_nuevos = productos_nuevos_count
            lote.productos_actualizados = productos_actualizados_costo
            lote.productos_sin_cambios = productos_sin_cambios
            lote.conflictos_detectados = conflictos_count
            lote.productos_sin_imagen = productos_sin_imagen
            lote.resumen_json = resumen
            if is_dry_run:
                lote.estado = "CON_CONFLICTOS" if conflictos_count > 0 else "SIMULADO"
            else:
                lote.estado = "APLICADO"
                lote.es_dry_run = False
            lote.save()

        # Auditoría si fue definitivo
        if not is_dry_run:
            registrar_actividad(
                request=request,
                accion="IMPORTACION_EJECUTADA",
                modelo_afectado="Producto",
                objeto_id=lote.id if lote else "",
                descripcion=f"Importación de catálogo ejecutada: {productos_nuevos_count} nuevos, {productos_actualizados_costo} costos actualizados.",
                detalles={
                    "nuevos": productos_nuevos_count,
                    "costos_actualizados": productos_actualizados_costo,
                    "sin_cambios": productos_sin_cambios,
                    "conflictos": conflictos_count,
                }
            )

        return resumen
