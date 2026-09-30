"""
Tests para ImportacionCatalogoService con soporte para lista comercial nueva y tramos de volumen.
"""
from decimal import Decimal
from pathlib import Path
import tempfile
import openpyxl

from django.test import TestCase

from apps.catalogo.models import Proveedor, Categoria, Producto, PrecioProveedorTramo
from apps.catalogo.services_importacion import ImportacionCatalogoService


class ImportacionNuevaListaTestCase(TestCase):
    def setUp(self):
        self.proveedor, _ = Proveedor.objects.get_or_create(
            codigo="KEY",
            defaults={"nombre": "Keyestudio", "moneda_origen": "USD"}
        )
        self.categoria, _ = Categoria.objects.get_or_create(
            nombre="Kits educativos",
            defaults={"descripcion": "Kits educativos"}
        )

    def _crear_excel_mock(self, filas):
        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        wb = openpyxl.Workbook()
        ws = wb.active
        for r in filas:
            ws.append(r)
        wb.save(tmp.name)
        wb.close()
        return tmp.name

    def test_detectar_formato_lista_comercial_nueva(self):
        filas = [
            ["HOT PRODUCTS", None, None, None, None, None, None, None],
            ["Image", "SKU", "Product Name", "Features", "Q1-9", "Q10-49", "Q50-100", "Q101-300"],
            [None, "KS0540", "Starter Kit", "Features text", 14.0, 13.5, 13.0, 12.5],
        ]
        path = self._crear_excel_mock(filas)
        wb = openpyxl.load_workbook(path, data_only=True)
        formato = ImportacionCatalogoService.detectar_formato(wb.active)
        wb.close()
        self.assertEqual(formato, ImportacionCatalogoService.FORMATO_LISTA_COMERCIAL_NUEVA)

    def test_detectar_formato_maestro_antiguo(self):
        filas = [
            ["SKU o ID", "Descripción del producto", "Miniatura", "Precio (USD)", "Imagen alta resolución"],
            ["KS0540", "Starter Kit Original", None, "26.50 USD", "Ver imagen"],
        ]
        path = self._crear_excel_mock(filas)
        wb = openpyxl.load_workbook(path, data_only=True)
        formato = ImportacionCatalogoService.detectar_formato(wb.active)
        wb.close()
        self.assertEqual(formato, ImportacionCatalogoService.FORMATO_MAESTRO_ANTIGUO)

    def test_validar_tramos_volumen_monotonia_correcta(self):
        q1_9 = Decimal("14.00")
        q10_49 = Decimal("13.50")
        q50_100 = Decimal("13.00")
        q101_300 = Decimal("12.50")

        tramos, tiene_anomalia, anomalias = ImportacionCatalogoService.validar_tramos_volumen(
            q1_9, q10_49, q50_100, q101_300
        )
        self.assertFalse(tiene_anomalia)
        self.assertEqual(len(anomalias), 0)
        self.assertEqual(len(tramos), 4)
        for t in tramos:
            self.assertEqual(t["estado_validacion"], "VALIDADO")
            self.assertFalse(t["es_anomalo"])

    def test_validar_tramos_volumen_anomalia_detectada(self):
        # Q101-300 mayor que Q50-100
        q1_9 = Decimal("16.00")
        q10_49 = Decimal("15.50")
        q50_100 = Decimal("15.00")
        q101_300 = Decimal("34.00")

        tramos, tiene_anomalia, anomalias = ImportacionCatalogoService.validar_tramos_volumen(
            q1_9, q10_49, q50_100, q101_300
        )
        self.assertTrue(tiene_anomalia)
        self.assertEqual(len(anomalias), 1)
        self.assertIn("Q101-300", anomalias[0])

        tramo_300 = next(t for t in tramos if t["cantidad_minima"] == 101)
        self.assertTrue(tramo_300["es_anomalo"])
        self.assertEqual(tramo_300["estado_validacion"], "PRECIO_PROVEEDOR_REQUIERE_REVISION")

    def test_dry_run_no_modifica_base_de_datos(self):
        prod = Producto.objects.create(
            sku_humm="HUMM-KEY-KS0540",
            sku_proveedor="KS0540",
            proveedor=self.proveedor,
            categoria=self.categoria,
            nombre_comercial="Kit Original Guardado",
            costo_proveedor_usd=Decimal("26.50"),
            publicado=False,
            activo=True,
            estado_curaduria="SIN_REVISAR"
        )

        filas = [
            ["HOT PRODUCTS", None, None, None, None, None, None, None],
            ["Image", "SKU", "Product Name", "Features", "Q1-9", "Q10-49", "Q50-100", "Q101-300"],
            [None, "KS0540", "Keyestudio Basic Starter Kit For Arduino With Board", "* 20 projects", 14.0, 13.5, 13.0, 12.5],
        ]
        path = self._crear_excel_mock(filas)

        res = ImportacionCatalogoService.procesar_catalogo(
            excel_path=path,
            is_dry_run=True,
            skus_filtro=["KS0540"]
        )

        # Verificar reporte
        self.assertTrue(res["is_dry_run"])
        self.assertEqual(res["productos_actualizados"], 1)
        self.assertEqual(len(res["detalle_skus"]), 1)
        self.assertEqual(res["detalle_skus"][0]["sku"], "KS0540")
        self.assertEqual(res["detalle_skus"][0]["q1_9"], 14.0)

        # Verificar que la BD NO cambió
        prod.refresh_from_db()
        self.assertEqual(prod.costo_proveedor_usd, Decimal("26.50"))
        self.assertEqual(PrecioProveedorTramo.objects.filter(producto=prod).count(), 0)

    def test_preservacion_curaduria_existente_en_ejecucion(self):
        prod = Producto.objects.create(
            sku_humm="HUMM-KEY-KS0530",
            sku_proveedor="KS0530",
            proveedor=self.proveedor,
            categoria=self.categoria,
            nombre_comercial="Kit Solar Curado Por Humm",
            descripcion_educativa="Descripción pedagógica protegida.",
            costo_proveedor_usd=Decimal("65.90"),
            publicado=True,
            activo=True,
            estado_curaduria="VALIDADO"
        )

        filas = [
            ["HOT PRODUCTS", None, None, None, None, None, None, None],
            ["Image", "SKU", "Product Name", "Features", "Q1-9", "Q10-49", "Q50-100", "Q101-300"],
            [None, "KS0530", "Keyestudio Solar Tracking Kit.", "* Engage hands-on", 37.0, 36.0, 35.0, 33.5],
        ]
        path = self._crear_excel_mock(filas)

        res = ImportacionCatalogoService.procesar_catalogo(
            excel_path=path,
            is_dry_run=False,
            skus_filtro=["KS0530"]
        )

        prod.refresh_from_db()
        # Costo actualizado
        self.assertEqual(prod.costo_proveedor_usd, Decimal("37.00"))
        # Features agregadas
        self.assertIn("Engage hands-on", prod.features_proveedor)
        # Curaduría Humm estrictamente preservada
        self.assertEqual(prod.nombre_comercial, "Kit Solar Curado Por Humm")
        self.assertEqual(prod.descripcion_educativa, "Descripción pedagógica protegida.")
        self.assertEqual(prod.estado_curaduria, "VALIDADO")
        self.assertTrue(prod.publicado)

        # Tramos de volumen registrados
        tramos = PrecioProveedorTramo.objects.filter(producto=prod).order_by("cantidad_minima")
        self.assertEqual(tramos.count(), 4)
        self.assertEqual(tramos[0].precio_usd, Decimal("37.00"))
        self.assertEqual(tramos[1].precio_usd, Decimal("36.00"))
        self.assertEqual(tramos[2].precio_usd, Decimal("35.00"))
        self.assertEqual(tramos[3].precio_usd, Decimal("33.50"))

    def test_creacion_nuevo_producto_con_tramos(self):
        filas = [
            ["HOT PRODUCTS", None, None, None, None, None, None, None],
            ["Image", "SKU", "Product Name", "Features", "Q1-9", "Q10-49", "Q50-100", "Q101-300"],
            [None, "FKS0003", "ESP32 Smart Robot Arm", "* 5 DOF robotic arm", 29.91, 28.91, 27.92, 26.92],
        ]
        path = self._crear_excel_mock(filas)

        res = ImportacionCatalogoService.procesar_catalogo(
            excel_path=path,
            is_dry_run=False,
            skus_filtro=["FKS0003"]
        )

        prod = Producto.objects.get(sku_proveedor="FKS0003")
        self.assertFalse(prod.publicado)
        self.assertTrue(prod.activo)
        self.assertEqual(prod.estado_curaduria, "SIN_REVISAR")
        self.assertEqual(prod.costo_proveedor_usd, Decimal("29.91"))
        self.assertIn("5 DOF", prod.features_proveedor)

        tramos = PrecioProveedorTramo.objects.filter(producto=prod).order_by("cantidad_minima")
        self.assertEqual(tramos.count(), 4)
        self.assertEqual(tramos[0].precio_usd, Decimal("29.91"))
        self.assertEqual(tramos[3].precio_usd, Decimal("26.92"))

    def test_dry_run_estrictamente_read_only_sin_dependencias(self):
        # Asegurar base limpia sin Proveedor, Categoría, Producto ni PrecioProveedorTramo
        Producto.objects.all().delete()
        PrecioProveedorTramo.objects.all().delete()
        Proveedor.objects.all().delete()
        Categoria.objects.all().delete()

        prov_count_inicial = Proveedor.objects.count()
        cat_count_inicial = Categoria.objects.count()
        prod_count_inicial = Producto.objects.count()
        tramos_count_inicial = PrecioProveedorTramo.objects.count()

        filas = [
            ["HOT PRODUCTS", None, None, None, None, None, None, None],
            ["Image", "SKU", "Product Name", "Features", "Q1-9", "Q10-49", "Q50-100", "Q101-300"],
            [None, "KS4050", "Environment Monitoring Learning Kit", "* Sensor kit", 39.88, 37.22, 35.89, 34.56],
        ]
        path = self._crear_excel_mock(filas)

        res = ImportacionCatalogoService.procesar_catalogo(
            excel_path=path,
            is_dry_run=True,
            skus_filtro=["KS4050"]
        )

        # 1. El reporte debe reflejar la simulación y reportar dependencias faltantes
        self.assertTrue(res["is_dry_run"])
        self.assertEqual(res["estado_dependencias"], "DEPENDENCIA_FALTANTE")
        self.assertTrue(len(res["dependencias_faltantes"]) >= 2)

        # 2. Las tablas de negocio NO deben haberse alterado en lo absoluto (100% READ-ONLY)
        self.assertEqual(Proveedor.objects.count(), prov_count_inicial)
        self.assertEqual(Categoria.objects.count(), cat_count_inicial)
        self.assertEqual(Producto.objects.count(), prod_count_inicial)
        self.assertEqual(PrecioProveedorTramo.objects.count(), tramos_count_inicial)

    def test_sincronizar_imagenes_banco(self):
        import tempfile
        from PIL import Image

        prov = Proveedor.objects.create(codigo="TST", nombre="Test", moneda_origen="USD")
        cat = Categoria.objects.create(nombre="Test Cat")
        prod = Producto.objects.create(
            proveedor=prov,
            categoria=cat,
            sku_humm="HUMM-TST-KS9999",
            sku_proveedor="KS9999",
            nombre_comercial="Producto Test",
            nombre_original_proveedor="Product Test Original",
            costo_proveedor_usd=Decimal("10.00"),
            precio_sugerido_total_clp=Decimal("20000"),
            publicado=False,
            activo=True,
        )

        with tempfile.TemporaryDirectory() as tmp_dir:
            img_path = Path(tmp_dir) / "KS9999.png"
            img = Image.new("RGBA", (100, 100), color=(255, 0, 0, 255))
            img.save(img_path)

            # 1. Primera sincronización -> vincula correctamente
            reporte = ImportacionCatalogoService.sincronizar_imagenes_banco(
                skus=["KS9999", "KS_NO_EXISTE"],
                imagenes_dir=tmp_dir
            )

            self.assertEqual(len(reporte), 2)
            rep_ks9999 = next(r for r in reporte if r["sku"] == "KS9999")
            self.assertTrue(rep_ks9999["imagen_vinculada"])
            self.assertEqual(rep_ks9999["estado"], "OK")
            self.assertEqual(rep_ks9999["archivo_encontrado"], "KS9999.png")

            # Verificar en BD
            self.assertEqual(prod.imagenes.count(), 1)
            img_obj = prod.imagenes.first()
            self.assertTrue(img_obj.es_principal)
            self.assertEqual(img_obj.nombre_archivo_original, "KS9999.png")

            # 2. Segunda sincronización -> detecta ya vinculada sin duplicar
            reporte2 = ImportacionCatalogoService.sincronizar_imagenes_banco(
                skus=["KS9999"],
                imagenes_dir=tmp_dir
            )
            self.assertFalse(reporte2[0]["imagen_vinculada"])
            self.assertEqual(reporte2[0]["estado"], "YA_VINCULADA")
            self.assertEqual(prod.imagenes.count(), 1)


